#!/usr/bin/env python3
"""
Excel Data Cleaning Script for MasterList.xlsx

This script cleans data in two ways:
1. Column D (Industry): Converts industry strings to standardized arrays using Groq LLM
2. Column E (LinkedIn URL): Extracts actual URLs from hyperlinks

Requirements:
- pip install pandas openpyxl groq python-dotenv
- Set VITE_GROQ_API_KEY in .env file
"""

import pandas as pd
import re
import json
import os
from groq import Groq
from typing import List, Optional
import time
from urllib.parse import urlparse
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv()

class ExcelDataCleaner:
    def __init__(self, groq_api_key: Optional[str] = None):
        """Initialize the data cleaner with Groq API key."""
        self.groq_api_key = groq_api_key or os.getenv('VITE_GROQ_API_KEY')
        if not self.groq_api_key:
            raise ValueError("Groq API key is required. Set VITE_GROQ_API_KEY in .env file or pass it directly.")
        
        self.client = Groq(api_key=self.groq_api_key)
        self.model = "openai/gpt-oss-20b"  # Using OpenAi as it's available on Groq
        
    def clean_industry_with_llm(self, industry_text: str) -> List[str]:
        """
        Use Groq LLM to convert industry text to standardized array format.
        
        Args:
            industry_text: Raw industry string like "Banking / FinTech" or "Other / Unknown"
            
        Returns:
            List of cleaned industry categories
        """
        if not industry_text or pd.isna(industry_text):
            return ["Unknown"]
        
        prompt = f"""
Convert the following industry text into a clean array of industry categories.

Rules:
1. Split on separators like "/", "&", ",", "|"
2. Remove extra whitespace
3. Standardize common terms (e.g., "FinTech" not "fintech", "AI" not "ai")
4. Return as a JSON array of strings
5. If unclear, keep original terms but clean them

Industry text: "{industry_text}"

Return only the JSON array, nothing else:
"""

        try:
            response = self.client.chat.completions.create(
                model=self.model,
                messages=[
                    {"role": "system", "content": "You are a data cleaning expert. Return only valid JSON arrays."},
                    {"role": "user", "content": prompt}
                ],
                temperature=0.1,
                max_tokens=200
            )
            
            result = response.choices[0].message.content.strip()
            
            # Parse the JSON response
            try:
                industry_array = json.loads(result)
                if isinstance(industry_array, list):
                    return [str(item).strip() for item in industry_array if str(item).strip()]
                else:
                    # Fallback to manual parsing
                    return self._manual_industry_split(industry_text)
            except json.JSONDecodeError:
                # Fallback to manual parsing
                return self._manual_industry_split(industry_text)
                
        except Exception as e:
            print(f"Error calling Groq API for '{industry_text}': {e}")
            # Fallback to manual parsing
            return self._manual_industry_split(industry_text)
    
    def _manual_industry_split(self, industry_text: str) -> List[str]:
        """
        Fallback method to manually split industry text.
        
        Args:
            industry_text: Raw industry string
            
        Returns:
            List of industry categories
        """
        if not industry_text or pd.isna(industry_text):
            return ["Unknown"]
        
        # Split on common separators
        separators = [' / ', '/', ' & ', '&', ', ', ',', ' | ', '|']
        parts = [industry_text]
        
        for sep in separators:
            new_parts = []
            for part in parts:
                new_parts.extend(part.split(sep))
            parts = new_parts
        
        # Clean and filter parts
        cleaned_parts = []
        for part in parts:
            cleaned = part.strip()
            if cleaned and cleaned.lower() not in ['and', 'or', 'the']:
                cleaned_parts.append(cleaned)
        
        return cleaned_parts if cleaned_parts else ["Unknown"]
    
    def extract_linkedin_url(self, cell_value, hyperlink_info=None) -> str:
        """
        Extract LinkedIn URL from cell value or hyperlink information.
        
        Args:
            cell_value: The visible text in the cell
            hyperlink_info: Hyperlink information if available
            
        Returns:
            Clean LinkedIn URL or empty string
        """
        if not cell_value and not hyperlink_info:
            return ""
        
        # First priority: Check hyperlink_info if available
        if hyperlink_info and isinstance(hyperlink_info, str):
            if 'linkedin.com' in hyperlink_info.lower():
                return self._clean_linkedin_url(hyperlink_info)
        
        # Second priority: Check if cell_value is already a valid URL
        if cell_value and isinstance(cell_value, str):
            if cell_value.startswith('http') and 'linkedin.com' in cell_value.lower():
                return self._clean_linkedin_url(cell_value)
        
        # If cell_value looks like "Link" or similar, but we have no hyperlink, return empty
        if cell_value and isinstance(cell_value, str):
            if cell_value.lower().strip() in ['link', 'linkedin', 'profile']:
                return ""  # Display text only, no actual URL available
        
        return ""
    
    def _clean_linkedin_url(self, url: str) -> str:
        """Clean and validate LinkedIn URL."""
        if not url:
            return ""
        
        url = url.strip()
        
        # Basic URL validation
        try:
            parsed = urlparse(url)
            if 'linkedin.com' in parsed.netloc:
                return url
        except:
            pass
        
        return ""
    
    def process_excel_file(self, input_file: str, output_file: str = None, batch_size: int = 50, task_choice: str = "both"):
        """
        Process the Excel file and clean the data.
        
        Args:
            input_file: Path to input Excel file
            output_file: Path to output Excel file (optional)
            batch_size: Number of rows to process at once for LLM calls
            task_choice: Which task to perform ("industry", "hyperlinks", or "both")
        """
        print(f"🔄 Loading Excel file: {input_file}")
        
        # Read the Excel file with pandas and also load workbook for hyperlinks
        try:
            df = pd.read_excel(input_file, engine='openpyxl')
            print(f"✅ Loaded {len(df)} rows from Excel file")
            
            # Also load the workbook directly to access hyperlinks
            from openpyxl import load_workbook
            wb = load_workbook(input_file)
            ws = wb.active
            print(f"✅ Loaded workbook for hyperlink extraction")
        except Exception as e:
            print(f"❌ Error reading Excel file: {e}")
            return
        
        print(f"📋 Columns: {list(df.columns)}")
        print(f"🎯 Task selected: {task_choice}")
        
        # Verify required columns exist based on task choice
        required_columns = []
        if task_choice in ["industry", "both"]:
            required_columns.append('Industry')
        if task_choice in ["hyperlinks", "both"]:
            required_columns.append('LinkedIn URL')
            
        missing_columns = [col for col in required_columns if col not in df.columns]
        if missing_columns:
            print(f"❌ Missing required columns: {missing_columns}")
            return
        
        # Create output columns based on task choice
        if task_choice in ["industry", "both"]:
            df['Industry_Cleaned'] = ""
        if task_choice in ["hyperlinks", "both"]:
            df['LinkedIn_URL_Cleaned'] = ""
        
        print("\n🔄 Processing data...")
        total_rows = len(df)
        
        # Process industries in batches to manage API rate limits (only if selected)
        if task_choice in ["industry", "both"]:
            print(f"📈 Processing industries with AI (this may take a while)...")
            processed_industries = 0
            
            for i in range(0, total_rows, batch_size):
                batch_end = min(i + batch_size, total_rows)
                print(f"Processing industries {i+1}-{batch_end} of {total_rows}...")
                
                for idx in range(i, batch_end):
                    industry = df.iloc[idx]['Industry']
                    cleaned_industry = self.clean_industry_with_llm(industry)
                    df.iloc[idx, df.columns.get_loc('Industry_Cleaned')] = json.dumps(cleaned_industry)
                    processed_industries += 1
                
                # Small delay to respect rate limits
                time.sleep(1)
            
            print(f"✅ Processed {processed_industries} industry records")
        else:
            print(f"⏭️ Skipping industry processing (not selected)")
        
        # Process LinkedIn URLs (only if selected)
        if task_choice in ["hyperlinks", "both"]:
            print("\n🔗 Processing LinkedIn URLs...")
            processed_urls = 0
            
            # Find the LinkedIn URL column index
            linkedin_col_index = None
            for col_idx, col_name in enumerate(df.columns):
                if col_name == 'LinkedIn URL':
                    linkedin_col_index = col_idx + 1  # openpyxl uses 1-based indexing
                    break
            
            for idx, row in df.iterrows():
                linkedin_cell = row['LinkedIn URL']
                
                # Try to extract hyperlink from the workbook
                hyperlink_url = ""
                if linkedin_col_index:
                    try:
                        cell = ws.cell(row=idx + 2, column=linkedin_col_index)  # +2 because pandas is 0-based, openpyxl is 1-based, and we have headers
                        if cell.hyperlink:
                            hyperlink_url = cell.hyperlink.target
                            if processed_urls < 10:  # Only show first 10 for brevity
                                print(f"Found hyperlink in row {idx + 2}: {hyperlink_url}")
                    except Exception as e:
                        pass  # Continue with regular processing
                
                cleaned_url = self.extract_linkedin_url(linkedin_cell, hyperlink_url)
                df.iloc[idx, df.columns.get_loc('LinkedIn_URL_Cleaned')] = cleaned_url
                processed_urls += 1
                
                # Show progress every 1000 rows
                if processed_urls % 1000 == 0:
                    print(f"Processed {processed_urls}/{total_rows} LinkedIn URLs...")
            
            print(f"✅ Processed {processed_urls} LinkedIn URL records")
        else:
            print(f"⏭️ Skipping LinkedIn URL processing (not selected)")
        
        # Generate output filename if not provided
        if not output_file:
            base_name = os.path.splitext(input_file)[0]
            output_file = f"{base_name}_cleaned.xlsx"
        
        # Save the cleaned data
        try:
            df.to_excel(output_file, index=False, engine='openpyxl')
            print(f"✅ Cleaned data saved to: {output_file}")
        except Exception as e:
            print(f"❌ Error saving file: {e}")
            return
        finally:
            # Close the workbook
            try:
                wb.close()
            except:
                pass
        
        # Print summary statistics
        self._print_summary(df, task_choice)
        
        return df
    
    def _print_summary(self, df: pd.DataFrame, task_choice: str = "both"):
        """Print summary statistics of the cleaning process."""
        print("\n📊 CLEANING SUMMARY:")
        print("=" * 50)
        
        # Industry statistics (only if industry cleaning was performed)
        if task_choice in ["industry", "both"] and 'Industry_Cleaned' in df.columns:
            industry_stats = {
                'total_industries': len(df),
                'empty_industries': df['Industry'].isna().sum(),
                'cleaned_industries': (~df['Industry_Cleaned'].str.contains('Unknown')).sum()
            }
            
            print(f"📈 Industry Cleaning:")
            print(f"   Total records: {industry_stats['total_industries']}")
            print(f"   Empty industries: {industry_stats['empty_industries']}")
            print(f"   Successfully cleaned: {industry_stats['cleaned_industries']}")
            
            # Sample cleaned industries
            print(f"\n🎯 Sample Cleaned Industries:")
            sample_industries = df[df['Industry_Cleaned'] != ''].head(5)
            for _, row in sample_industries.iterrows():
                original = row['Industry']
                cleaned = row['Industry_Cleaned']
                print(f"   '{original}' → {cleaned}")
        
        # LinkedIn URL statistics (only if hyperlink extraction was performed)
        if task_choice in ["hyperlinks", "both"] and 'LinkedIn_URL_Cleaned' in df.columns:
            url_stats = {
                'total_urls': len(df),
                'empty_original': df['LinkedIn URL'].isna().sum(),
                'valid_cleaned': (df['LinkedIn_URL_Cleaned'] != "").sum(),
                'empty_cleaned': (df['LinkedIn_URL_Cleaned'] == "").sum()
            }
            
            print(f"\n🔗 LinkedIn URL Cleaning:")
            print(f"   Total records: {url_stats['total_urls']}")
            print(f"   Empty original URLs: {url_stats['empty_original']}")
            print(f"   Valid cleaned URLs: {url_stats['valid_cleaned']}")
            print(f"   Empty after cleaning: {url_stats['empty_cleaned']}")
            
            # Sample cleaned URLs
            print(f"\n🎯 Sample Cleaned URLs:")
            sample_urls = df[df['LinkedIn_URL_Cleaned'] != ''].head(5)
            for _, row in sample_urls.iterrows():
                original = row['LinkedIn URL']
                cleaned = row['LinkedIn_URL_Cleaned']
                print(f"   '{original}' → {cleaned}")

def get_task_choice():
    """Get user's choice of which task to perform."""
    print("📋 SELECT TASK TO PERFORM:")
    print("=" * 40)
    print("1. Industry Cleaning (AI-powered, slow)")
    print("   - Converts industry text to standardized arrays")
    print("   - Uses Groq LLM for intelligent parsing")
    print()
    print("2. Hyperlink Extraction (Fast)")
    print("   - Extracts LinkedIn URLs from hyperlinks")
    print("   - Processes 'Link' cells to actual URLs")
    print()
    print("3. Both Tasks (Slowest)")
    print("   - Performs both industry cleaning and hyperlink extraction")
    print()
    
    while True:
        choice = input("Enter your choice (1, 2, or 3): ").strip()
        if choice == "1":
            return "industry"
        elif choice == "2":
            return "hyperlinks"
        elif choice == "3":
            return "both"
        else:
            print("❌ Invalid choice. Please enter 1, 2, or 3.")

def main():
    """Main function to run the data cleaning script."""
    # Configuration
    INPUT_FILE = r"C:\Users\hweth\OneDrive\Desktop\Innovera\Rapid Prototyping\frank\data_clean\MasterList.xlsx"
    OUTPUT_FILE = None  # Will auto-generate if None
    GROQ_API_KEY = None  # Will use environment variable from .env file
    
    print("🚀 EXCEL DATA CLEANING SCRIPT")
    print("=" * 50)
    
    # Get user's task choice
    task_choice = get_task_choice()
    
    print(f"\n🎯 Selected task: {task_choice.upper()}")
    print("=" * 50)
    
    try:
        # Initialize the cleaner (only if needed for industry processing)
        if task_choice in ["industry", "both"]:
            cleaner = ExcelDataCleaner(groq_api_key=GROQ_API_KEY)
        else:
            # For hyperlink-only processing, we don't need Groq API
            cleaner = ExcelDataCleaner.__new__(ExcelDataCleaner)
            cleaner.groq_api_key = None
            cleaner.client = None
            cleaner.model = None
        
        # Process the file
        cleaned_df = cleaner.process_excel_file(
            input_file=INPUT_FILE,
            output_file=OUTPUT_FILE,
            batch_size=25,  # Adjust based on your API rate limits
            task_choice=task_choice
        )
        
        if cleaned_df is not None:
            print("\n✨ Data cleaning completed successfully!")
            print(f"📁 Check your output file for the cleaned data.")
        
    except Exception as e:
        print(f"❌ Error: {e}")
        print("\nMake sure to:")
        print("1. Install required packages: pip install pandas openpyxl groq python-dotenv")
        print("2. Set your Groq API key in the .env file: VITE_GROQ_API_KEY='your-api-key'")
        print("3. Ensure the Excel file exists at the specified path")

if __name__ == "__main__":
    main()